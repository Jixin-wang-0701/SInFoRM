"""Verify packaged matrices, coefficients, score CSVs and Excel rows."""

from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy import sparse


ROOT = Path(__file__).resolve().parents[1]
MATRICES = ROOT / "matrices"
COEF = ROOT / "coefficients"
SCORES = ROOT / "scores"


def near(a: float, b: float) -> bool:
    return math.isclose(a, b, rel_tol=1e-11, abs_tol=1e-12)


def main() -> None:
    rows = pd.read_csv(MATRICES / "rows.csv", usecols=["bed_id", "Chromosome", "ZSTAT"])
    assert len(rows) == 95_750 and rows.bed_id.is_unique
    full = pd.read_csv(SCORES / "full_fit_scores.csv")
    loco = pd.read_csv(SCORES / "loco_scores.csv")
    for table in (full, loco):
        assert len(table) == len(rows) and table.bed_id.equals(rows.bed_id)
        assert table.Chromosome.equals(rows.Chromosome)
        assert np.allclose(table.ZSTAT, rows.ZSTAT)
        assert np.allclose(table.signed_sum,
                           table.expression_component + table.go_component + table.pfam_component)
        assert np.allclose(table.feature_only_abs_score, table.signed_sum.abs())
        assert sorted(table.rank_desc.tolist()) == list(range(1, len(rows) + 1))

    for block, coefficient_file in (("expression", "coef_expression.npy"),
                                    ("go", "coef_go.npy"),
                                    ("pfam", "coef_pfam_new.npy")):
        matrix_file = "X_pfam_rebuilt.npz" if block == "pfam" else f"X_{block}.npz"
        matrix = sparse.load_npz(MATRICES / matrix_file)
        coefficients = np.load(COEF / coefficient_file)
        calculated = np.asarray(matrix @ coefficients).ravel()
        assert np.allclose(calculated, full[f"{block}_component"], atol=1e-12)
        print(block, matrix.shape, matrix.nnz)

    prior = pd.read_csv(ROOT / "full_fit" / "scores_comparison.csv",
                        usecols=["bed_id", "all_junction_refit_new_pfam"])
    assert prior.bed_id.equals(rows.bed_id)
    assert np.allclose(full.feature_only_abs_score, prior.all_junction_refit_new_pfam,
                       atol=1e-12)
    prior_loco = pd.read_csv(ROOT / "loco" / "oof_predictions.csv")
    assert prior_loco.bed_id.equals(rows.bed_id)
    assert np.allclose(loco.feature_only_abs_score,
                       prior_loco.rebuilt_feature_only_score, atol=1e-12)

    workbook = load_workbook(SCORES / "junction_scores_full_fit_and_loco.xlsx",
                             read_only=True, data_only=True)
    for sheet_name, filename in (("Full fit", "full_fit_scores.csv"),
                                 ("LOCO", "loco_scores.csv")):
        sheet = workbook[sheet_name]
        assert sheet.max_row == 95_751 and sheet.max_column == 12
        with (SCORES / filename).open(newline="", encoding="utf-8") as handle:
            reader = csv.reader(handle)
            sheet_rows = sheet.iter_rows(values_only=True)
            assert tuple(next(reader)) == tuple(next(sheet_rows))
            for n, (csv_row, xlsx_row) in enumerate(zip(reader, sheet_rows), start=1):
                assert csv_row[0:2] == list(xlsx_row[0:2]), (sheet_name, n)
                assert all(near(float(csv_row[i]), float(xlsx_row[i]))
                           for i in range(2, 12)), (sheet_name, n)
            assert n == 95_750, (sheet_name, n)
        print(sheet_name, n, "Excel rows verified")
    assert "ReadMe" in workbook.sheetnames
    workbook.close()
    print("All release checks passed")


if __name__ == "__main__":
    main()
